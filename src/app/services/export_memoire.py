"""Service — export téléchargeable du mémoire technique (lot L6, phase 4).

Objectif : **un mémoire qu'on ne peut pas sortir du site ne sert à rien**. Ce module
transforme le mémoire validé en fichier ouvrable dans un traitement de texte, sans
jamais franchir la ligne rouge :

* le fichier porte le contenu **validé** — l'empreinte SHA-256 enregistrée dans
  `memoire_validation` est **recalculée** et doit correspondre, sinon l'export est
  **refusé** (une relecture est exigée : on ne diffuse pas autre chose que ce qu'un
  humain nommé a relu) ;
* **rien ne sort sans validation humaine** : sans ligne `memoire_validation`, l'export
  n'est pas une erreur technique mais un message qui dit quoi faire ;
* le **Markdown (`.md`) est le format canonique**, produit **sans aucune dépendance
  nouvelle** (bibliothèque standard seule) : c'est lui qui est testé et vérifié ;
* le **DOCX est un format de confort** produit par `python-docx` (dépendance annoncée et
  justifiée dans `docs/MEMOIRE-TECHNIQUE.md` **avant** installation). Si la
  bibliothèque est absente, le lot **n'est pas bloqué** : ce format répond « non
  disponible » avec sa raison, et le Markdown reste valide à lui seul ;
* **aucun prix, aucun chiffre inventé, aucune conformité garantie** dans le fichier :
  le rendu ne fait que reprendre le contenu du mémoire (qui porte déjà cette règle) ;
* **isolation par client** : tout le SQL passe par `storage.connexion.Connexion` et le
  `client_id` de session. Aucun mémoire, aucune source d'un autre client.

Choix documentés dans `docs/MEMOIRE-TECHNIQUE.md` § 8 « Export téléchargeable ».
"""

from __future__ import annotations

import datetime as _dt
import io
import logging
import re
import unicodedata
from typing import Any, Mapping, Optional

from app.domain.memoire_technique_genere import AVERTISSEMENT_GENERATION
from app.services import memoire_technique
from app.storage.connexion import Connexion, ContexteClient

#: Journal du module.
LOGGER = logging.getLogger(__name__)

#: Format canonique : produit par la bibliothèque standard, testé et vérifié.
FORMAT_CANONIQUE = "md"

#: Formats d'export admis, dans l'ordre de préférence (le confort en dernier).
FORMATS_EXPORT: tuple[str, ...] = ("md", "docx")

#: Extension de fichier par format.
EXTENSIONS: dict[str, str] = {"md": "md", "docx": "docx"}

#: Type MIME servi par format (le `.md` est du texte UTF-8).
TYPES_MIME: dict[str, str] = {
    "md": "text/markdown; charset=utf-8",
    "docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ),
}

#: Longueur maximale du radical du nom de fichier (le reste est du slug).
LONGUEUR_SLUG = 60

#: Messages de refus — **pédagogiques**, jamais des erreurs techniques.

#: Refus quand aucune validation nommée et horodatée n'existe pour ce mémoire.
MESSAGE_VALIDATION_MANQUANTE = (
    "Le mémoire doit d'abord être relu et validé par une personne nommée. "
    "Rien ne sort sans relecture humaine : c'est la règle du produit. "
    "Pour débloquer le téléchargement : ouvrez le mémoire, relisez chaque section "
    "(une section passe de « à relire » à « relue »), puis validez le mémoire en "
    "indiquant votre nom et votre fonction. L'export devient disponible dans la foulée."
)

#: Refus quand le contenu a changé depuis la validation (empreinte divergente).
MESSAGE_CONTENU_DIVERGENT = (
    "Le contenu du mémoire a changé depuis sa validation : l'empreinte enregistrée "
    "ne correspond plus au contenu actuel. Par prudence, l'export est refusé — le "
    "fichier téléchargé doit porter exactement ce qu'une personne nommée a relu. "
    "Relisez les sections modifiées, puis validez à nouveau le mémoire."
)

#: Refus quand le format DOCX est demandé alors que `python-docx` est absent.
MESSAGE_DOCX_INDISPONIBLE = (
    "Le format Word (.docx) n'est pas disponible sur cette installation : {raison}. "
    "Le format Markdown (.md) reste disponible et suffit : il porte le même contenu "
    "validé et s'ouvre dans un traitement de texte."
)


# --------------------------------------------------------------------------- #
# Erreurs
# --------------------------------------------------------------------------- #
class ErreurExport(RuntimeError):
    """L'export ne peut pas aboutir — toujours par un refus explicite et lisible."""


class FormatInconnu(ErreurExport):
    """Le format demandé n'est pas un format d'export (refus explicite, liste des admis)."""


class ValidationManquante(ErreurExport):
    """Aucune validation humaine nommée et horodatée : l'export est refusé."""


class ContenuDivergent(ErreurExport):
    """L'empreinte du contenu actuel ne correspond pas à celle validée : export refusé."""


class FormatNonDisponible(ErreurExport):
    """Le format demandé exige une dépendance absente ; le canonique reste disponible."""


# --------------------------------------------------------------------------- #
# Format demandé
# --------------------------------------------------------------------------- #
def format_demande(valeur: object) -> str:
    """Normalise le paramètre `format` et refuse tout format inconnu.

    Une valeur absente vaut le format canonique (Markdown). Un format inconnu est un
    **refus explicite** qui rappelle la liste des formats admis — jamais un défaut
    silencieux.
    """
    brut = str(valeur or "").strip().casefold().lstrip(".")
    if not brut:
        return FORMAT_CANONIQUE
    if brut not in FORMATS_EXPORT:
        raise FormatInconnu(
            f"Format d'export inconnu : {brut!r}. Formats admis : "
            f"{', '.join(FORMATS_EXPORT)} (Markdown = format canonique)."
        )
    return brut


def docx_disponible() -> tuple[bool, str]:
    """Disponibilité de `python-docx`, avec sa raison lisible quand elle manque.

    L'import est fait **ici**, jamais au chargement du module : le Markdown doit rester
    produit même si la dépendance de confort est absente.
    """
    try:
        import docx  # noqa: F401 — simple sonde de disponibilité
    except Exception as exc:  # pragma: no cover — dépend de l'installation locale
        return False, f"bibliothèque python-docx absente ({type(exc).__name__})"
    return True, ""


# --------------------------------------------------------------------------- #
# Empreinte du contenu validé
# --------------------------------------------------------------------------- #
def _lire_pour_empreinte(
    connexion: Connexion, contexte: ContexteClient, dossier_id: str
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Relit sections (avec sources) et manques **dans l'ordre exact de la validation**.

    `memoire_technique.valider_dossier` calcule l'empreinte sur les sections triées par
    `ordre, id` et les manques triés par `id`. L'ordre compte : on recalcule donc avec
    les **mêmes tris**, sinon deux lectures identiques donneraient deux empreintes
    différentes.
    """
    sections = connexion.executer(
        contexte,
        "SELECT * FROM memoire_section WHERE client_id = %(client_id)s "
        "AND memoire_dossier_id = %(dossier)s ORDER BY ordre, id;",
        {"dossier": dossier_id},
    )
    for section in sections:
        section["sources"] = connexion.executer(
            contexte,
            "SELECT * FROM memoire_section_source WHERE client_id = %(client_id)s "
            "AND memoire_section_id = %(section)s ORDER BY date_creation, id;",
            {"section": str(section["id"])},
        )
    manques = connexion.executer(
        contexte,
        "SELECT * FROM memoire_manque WHERE client_id = %(client_id)s "
        "AND memoire_dossier_id = %(dossier)s ORDER BY id;",
        {"dossier": dossier_id},
    )
    return sections, manques


def empreinte_actuelle(
    connexion: Connexion, contexte: ContexteClient, dossier: Mapping[str, Any]
) -> str:
    """Empreinte SHA-256 du contenu **actuel** du mémoire, dans l'ordre de la validation."""
    sections, manques = _lire_pour_empreinte(connexion, contexte, str(dossier["id"]))
    return memoire_technique.empreinte_contenu(dossier, sections, manques)


# --------------------------------------------------------------------------- #
# Mise en forme lisible — jamais un identifiant interne
# --------------------------------------------------------------------------- #
def _poids_texte(poids: Any) -> str:
    """Pondération affichée sans zéro inutile (`40`, pas `40.00`)."""
    texte = str(poids)
    try:
        from decimal import Decimal

        texte = format(Decimal(texte).normalize(), "f")
    except Exception:  # pragma: no cover — valeur déjà textuelle
        pass
    return texte


def _horodatage_texte(valeur: Any) -> str:
    """Horodatage lisible en français, sans dépendre de la locale du poste."""
    if isinstance(valeur, (_dt.datetime, _dt.date)):
        if isinstance(valeur, _dt.datetime):
            return valeur.strftime("%d/%m/%Y à %H:%M")
        return valeur.strftime("%d/%m/%Y")
    return str(valeur or "").strip()


def _slug(texte: str) -> str:
    """Radical de nom de fichier : accents retirés, séparateurs normalisés, borné."""
    brut = unicodedata.normalize("NFKD", str(texte or ""))
    sans_accents = "".join(c for c in brut if not unicodedata.combining(c))
    propre = re.sub(r"[^A-Za-z0-9]+", "-", sans_accents).strip("-").casefold()
    return (propre[:LONGUEUR_SLUG].strip("-") or "")


def nom_fichier_export(dossier: Mapping[str, Any], format: str) -> str:
    """Nom du fichier téléchargé : lisible, **sans identifiant interne**.

    Le nom est tiré du titre du mémoire (une donnée métier), jamais d'un UUID : un
    fichier ne doit pas exposer de valeur technique. Un titre qui commence déjà par
    « Mémoire technique » n'est pas préfixé deux fois.
    """
    radical = re.sub(r"^(memoire-technique-)+", "", _slug(str(dossier.get("titre") or "")))
    base = f"memoire-technique-{radical}" if radical else "memoire-technique"
    return f"{base}.{EXTENSIONS[format]}"


def _corps_section(section: Mapping[str, Any]) -> str:
    """Contenu d'une section, sans la ligne de cadrage du critère déjà portée par le titre.

    Le générateur ouvre le contenu par « Critère « … » — pondération N %. » ; le rendu
    porte cette information dans le titre de section. On évite donc la répétition, sans
    retirer une seule phrase métier.
    """
    lignes = str(section.get("contenu") or "").strip().splitlines()
    if lignes and lignes[0].strip().startswith("Critère «"):
        lignes = lignes[1:]
    while lignes and not lignes[0].strip():
        lignes = lignes[1:]
    return "\n".join(lignes).strip()


# --------------------------------------------------------------------------- #
# Rendu Markdown — format canonique, bibliothèque standard seule
# --------------------------------------------------------------------------- #
def rendre_markdown(memoire: Mapping[str, Any]) -> str:
    """Rend le mémoire en Markdown : titres hiérarchisés, sources citées, manques listés.

    Aucune dépendance : le texte est assemblé à partir des seules valeurs du mémoire
    (titre, critère, pondération, contenu, sources, manques). Aucun chiffre n'est
    ajouté, aucun prix n'est fixé, aucune conformité n'est promise.
    """
    dossier = memoire.get("dossier") or {}
    sections = list(memoire.get("sections") or [])
    manques = list(memoire.get("manques") or [])
    validations = list(memoire.get("validations") or [])

    lignes: list[str] = []
    lignes.append(f"# {dossier.get('titre') or 'Mémoire technique'}")
    lignes.append("")
    lignes.append(f"*{memoire.get('avertissement') or AVERTISSEMENT_GENERATION}*")
    lignes.append("")
    lignes.append(
        "*Ce document est produit à partir de votre bibliothèque d'entreprise. Aucun "
        "prix n'est fixé ici, aucune conformité n'est garantie.*"
    )
    lignes.append("")

    if validations:
        validation = validations[0]
        lignes.append("## Relecture et validation humaine")
        lignes.append("")
        lignes.append(
            f"- Relu et validé par : {validation.get('nom_validateur')} "
            f"({validation.get('fonction_validateur')})"
        )
        lignes.append(
            f"- Date de validation : {_horodatage_texte(validation.get('horodatage'))}"
        )
        lignes.append(
            f"- Empreinte du contenu validé (SHA-256) : "
            f"{validation.get('empreinte_contenu')}"
        )
        lignes.append("")

    for index, section in enumerate(sections, start=1):
        titre = section.get("titre") or section.get("critere_libelle") or "Section"
        lignes.append(f"## {index}. {titre}")
        lignes.append("")
        poids = section.get("critere_poids")
        if poids is not None:
            lignes.append(
                f"*Critère « {section.get('critere_libelle')} » — pondération "
                f"{_poids_texte(poids)} %.*"
            )
        else:
            lignes.append(
                f"*Critère « {section.get('critere_libelle')} » — pondération non connue "
                "du DCE (traité en fin de mémoire).*"
            )
        lignes.append("")
        lignes.append(_corps_section(section))
        sources = list(section.get("sources") or [])
        if sources:
            lignes.append("")
            lignes.append("Sources citées :")
            lignes.append("")
            for source in sources:
                lignes.append(
                    f"- {source.get('libelle_source')} — "
                    f"{source.get('emplacement_source')}"
                )
        lignes.append("")

    lignes.append("## Manques à traiter")
    lignes.append("")
    if not manques:
        lignes.append(
            "Aucun manque signalé : tous les critères du DCE sont étayés par un élément "
            "réel de votre bibliothèque."
        )
        lignes.append("")
    else:
        lignes.append(
            "*Ces points ne sont pas des erreurs : ils disent ce que votre bibliothèque "
            "ne permet pas encore d'étayer, et l'action à mener pour y répondre.*"
        )
        lignes.append("")
        for manque in manques:
            poids = manque.get("critere_poids")
            suffixe = f" — pondération {_poids_texte(poids)} %" if poids is not None else ""
            lignes.append(f"### Critère « {manque.get('critere_libelle')} »{suffixe}")
            lignes.append("")
            lignes.append(f"- Constat : {manque.get('constat')}")
            lignes.append(f"- Action à mener : {manque.get('action_attendue')}")
            lignes.append("")

    return "\n".join(lignes).rstrip() + "\n"


# --------------------------------------------------------------------------- #
# Rendu DOCX — format de confort, dégradation propre si absent
# --------------------------------------------------------------------------- #
def _ajouter_paragraphes(document: Any, texte: str) -> None:
    """Écrit un contenu multi-lignes : les lignes `- …` deviennent des puces."""
    for ligne in str(texte or "").splitlines():
        morceau = ligne.strip()
        if not morceau:
            continue
        if morceau.startswith(("- ", "• ")):
            document.add_paragraph(morceau[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(morceau)


def rendre_docx(memoire: Mapping[str, Any]) -> bytes:
    """Rend le mémoire en `.docx` (binaire), via `python-docx`.

    Même contenu que le Markdown : titre, validation humaine, sections (titres
    hiérarchisés dans l'ordre des critères pondérés), sources citées, manques et actions.
    Lève `FormatNonDisponible` si la dépendance est absente — jamais un fichier partiel.
    """
    disponible, raison = docx_disponible()
    if not disponible:
        raise FormatNonDisponible(MESSAGE_DOCX_INDISPONIBLE.format(raison=raison))

    from docx import Document  # import tardif : le Markdown ne dépend pas de ce module

    dossier = memoire.get("dossier") or {}
    sections = list(memoire.get("sections") or [])
    manques = list(memoire.get("manques") or [])
    validations = list(memoire.get("validations") or [])

    document = Document()
    document.add_heading(str(dossier.get("titre") or "Mémoire technique"), level=0)

    note = document.add_paragraph(
        str(memoire.get("avertissement") or AVERTISSEMENT_GENERATION)
    )
    for run in note.runs:  # pragma: no branch — un seul run à la création
        run.italic = True
    document.add_paragraph(
        "Ce document est produit à partir de votre bibliothèque d'entreprise. "
        "Aucun prix n'est fixé ici, aucune conformité n'est garantie."
    )

    if validations:
        validation = validations[0]
        document.add_heading("Relecture et validation humaine", level=1)
        document.add_paragraph(
            f"Relu et validé par : {validation.get('nom_validateur')} "
            f"({validation.get('fonction_validateur')})"
        )
        document.add_paragraph(
            f"Date de validation : {_horodatage_texte(validation.get('horodatage'))}"
        )
        document.add_paragraph(
            f"Empreinte du contenu validé (SHA-256) : {validation.get('empreinte_contenu')}"
        )

    for index, section in enumerate(sections, start=1):
        titre = section.get("titre") or section.get("critere_libelle") or "Section"
        document.add_heading(f"{index}. {titre}", level=1)
        poids = section.get("critere_poids")
        if poids is not None:
            document.add_paragraph(
                f"Critère « {section.get('critere_libelle')} » — pondération "
                f"{_poids_texte(poids)} %."
            )
        else:
            document.add_paragraph(
                f"Critère « {section.get('critere_libelle')} » — pondération non connue "
                "du DCE (traité en fin de mémoire)."
            )
        _ajouter_paragraphes(document, _corps_section(section))
        sources = list(section.get("sources") or [])
        if sources:
            document.add_paragraph("Sources citées :")
            for source in sources:
                document.add_paragraph(
                    f"{source.get('libelle_source')} — {source.get('emplacement_source')}",
                    style="List Bullet",
                )

    document.add_heading("Manques à traiter", level=1)
    if not manques:
        document.add_paragraph(
            "Aucun manque signalé : tous les critères du DCE sont étayés par un élément "
            "réel de votre bibliothèque."
        )
    else:
        document.add_paragraph(
            "Ces points ne sont pas des erreurs : ils disent ce que votre bibliothèque ne "
            "permet pas encore d'étayer, et l'action à mener pour y répondre."
        )
        for manque in manques:
            poids = manque.get("critere_poids")
            suffixe = f" — pondération {_poids_texte(poids)} %" if poids is not None else ""
            document.add_heading(
                f"Critère « {manque.get('critere_libelle')} »{suffixe}", level=2
            )
            document.add_paragraph(
                f"Constat : {manque.get('constat')}", style="List Bullet"
            )
            document.add_paragraph(
                f"Action à mener : {manque.get('action_attendue')}", style="List Bullet"
            )

    tampon = io.BytesIO()
    document.save(tampon)
    return tampon.getvalue()


# --------------------------------------------------------------------------- #
# Préparation de l'export — la porte unique
# --------------------------------------------------------------------------- #
def preparer_export(
    connexion: Connexion,
    contexte: ContexteClient,
    *,
    consultation_id: str,
    format: Optional[str] = None,
) -> dict[str, Any]:
    """Prépare le fichier à télécharger pour le client de la session.

    Ordre des vérifications — chacune est un refus explicite, jamais un fichier partiel :

    1. le mémoire doit exister **pour ce client** (`MemoireIntrouvable` → 404) ;
    2. une ligne `memoire_validation` (nom, fonction, horodatage) doit exister, sinon
       `ValidationManquante` (message pédagogique) ;
    3. l'empreinte SHA-256 du contenu **actuel** doit égaler celle enregistrée, sinon
       `ContenuDivergent` (une relecture est exigée) ;
    4. le format doit être disponible (`FormatNonDisponible` pour un DOCX sans
       `python-docx`).

    Renvoie `format`, `nom_fichier`, `contenu` (octets), `type_mime`, l'empreinte
    vérifiée et la ligne de validation retenue.
    """
    format_effectif = format_demande(format)
    memoire = memoire_technique.lire_memoire(connexion, contexte, consultation_id)
    dossier = memoire["dossier"]

    validations = list(memoire.get("validations") or [])
    if not validations:
        raise ValidationManquante(MESSAGE_VALIDATION_MANQUANTE)
    validation = validations[0]  # la plus récente (horodatage décroissant)

    # Le fichier doit porter ce qui a été validé : on le vérifie, on ne le suppose pas.
    calculee = empreinte_actuelle(connexion, contexte, dossier)
    if str(calculee) != str(validation.get("empreinte_contenu")):
        LOGGER.warning(
            "Export refusé : empreinte divergente (dossier=%s, validée=%s, actuelle=%s)",
            dossier.get("id"),
            validation.get("empreinte_contenu"),
            calculee,
        )
        raise ContenuDivergent(MESSAGE_CONTENU_DIVERGENT)

    if format_effectif == FORMAT_CANONIQUE:
        contenu = rendre_markdown(memoire).encode("utf-8")
    else:
        contenu = rendre_docx(memoire)

    return {
        "format": format_effectif,
        "nom_fichier": nom_fichier_export(dossier, format_effectif),
        "contenu": contenu,
        "type_mime": TYPES_MIME[format_effectif],
        "empreinte_contenu": calculee,
        "dossier": dossier,
        "validation": validation,
    }


__all__ = [
    "EXTENSIONS",
    "FORMATS_EXPORT",
    "FORMAT_CANONIQUE",
    "MESSAGE_CONTENU_DIVERGENT",
    "MESSAGE_DOCX_INDISPONIBLE",
    "MESSAGE_VALIDATION_MANQUANTE",
    "TYPES_MIME",
    "ContenuDivergent",
    "ErreurExport",
    "FormatInconnu",
    "FormatNonDisponible",
    "ValidationManquante",
    "docx_disponible",
    "empreinte_actuelle",
    "format_demande",
    "nom_fichier_export",
    "preparer_export",
    "rendre_docx",
    "rendre_markdown",
]
